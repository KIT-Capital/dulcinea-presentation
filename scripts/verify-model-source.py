"""Read-only reconciliation against a private workbook; never copies or uploads it.

Usage: python scripts/verify-model-source.py --source /private/path/to/model.xlsx
Requires openpyxl. Build the site and run verify-financials.mjs separately to check
the displayed statements, margins, cash reconciliation and EN/ES/FR headlines.
"""
import argparse
from datetime import date, datetime
import hashlib
import json
import math
from pathlib import Path
import re
import warnings
import xml.etree.ElementTree as ET
import zipfile

import openpyxl

# This is a read-only pass; no save operation can remove workbook formatting.
warnings.filterwarnings('ignore', message='Conditional Formatting extension is not supported and will be removed')

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--source', type=Path, required=True)
args = parser.parse_args()
root = Path(__file__).resolve().parent.parent
source_hash = hashlib.sha256(args.source.read_bytes()).hexdigest()
financials = json.loads((root / 'content/financials.json').read_text(encoding='utf8'))
summary = json.loads((root / 'content/model-summary.json').read_text(encoding='utf8'))
properties = json.loads((root / 'content/properties.json').read_text(encoding='utf8'))
assert source_hash == financials['source']['sha256'] == summary['provenance']['sha256'], 'Wrong source workbook'
assert all(item['source']['sha256'] == source_hash for item in properties), 'Property source differs'

book = openpyxl.load_workbook(args.source, data_only=True, read_only=True)
formulas = openpyxl.load_workbook(args.source, data_only=False, read_only=True)
# Read each worksheet once; repeated random reads in read-only mode rescan XML.
values = {sheet.title: {cell.coordinate: cell.value for row in sheet for cell in row if cell.value is not None} for sheet in book}
formula_cells = {sheet.title: {cell.coordinate: getattr(cell.value, 'text', cell.value) for row in sheet for cell in row if cell.data_type == 'f'} for sheet in formulas}

def same(actual, expected, label):
    if isinstance(actual, (datetime, date)):
        actual = actual.isoformat()
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        assert isinstance(actual, (int, float)) and math.isfinite(actual), f'Missing numeric value: {label}'
        assert abs(actual - expected) < 1e-8, f'Value changed: {label}: {actual} != {expected}'
    else:
        assert actual == expected, f'Value changed: {label}'

financial_count = 0
for section in ['profitAndLoss', 'afterCarry', 'cashFlows', 'balanceSheet']:
    for row in financials[section]['rows']:
        sources = row.get('sourceCells', [])
        for index, cell in enumerate(sources):
            match = re.fullmatch(r"'([^']+)'!([A-Z]+[0-9]+)", cell['cell'])
            assert match, cell['cell']
            actual = values[match[1]].get(match[2])
            same(actual, cell['cachedValue'], cell['cell'])
            assert formula_cells[match[1]].get(match[2]) == cell['formula'], f"Formula changed: {cell['cell']}"
            if 'sourceSignMultiplier' in row:
                same(actual * row['sourceSignMultiplier'], row['values'][index], cell['cell'])
            financial_count += 1

summary_count = 0
def check_summary(node):
    global summary_count
    if isinstance(node, list):
        for value in node:
            check_summary(value)
    elif isinstance(node, dict):
        if 'value' in node and 'ref' in node:
            ref = node['ref']
            single = re.fullmatch(r"(-?)'([^']+)'!([A-Z]+[0-9]+)", ref)
            total = re.fullmatch(r"SUM\('([^']+)'!([A-Z]+)([0-9]+):([A-Z]+)([0-9]+)\)", ref)
            if single:
                actual = values[single[2]].get(single[3])
                if single[1]:
                    assert isinstance(actual, (int, float)), ref
                    actual = -actual
            elif total:
                min_col = openpyxl.utils.column_index_from_string(total[2])
                max_col = openpyxl.utils.column_index_from_string(total[4])
                cells = [values[total[1]].get(f'{openpyxl.utils.get_column_letter(col)}{row}')
                         for row in range(int(total[3]), int(total[5]) + 1)
                         for col in range(min_col, max_col + 1)]
                assert all(isinstance(value, (int, float)) for value in cells), f'Missing numeric cells: {ref}'
                actual = sum(cells)
            else:
                raise AssertionError(f'Unsupported reference: {ref}')
            same(actual, node['value'], ref)
            summary_count += 1
        for value in node.values():
            check_summary(value)
check_summary(summary)

for flow in summary['monthlyInvestorCashFlows']:
    match = re.fullmatch(r"'_IRR Engine'!B(\d+):G\1", flow['ref'])
    assert match, flow['ref']
    row = match[1]
    for key, column, sign in [('date', 'B', 1), ('calledUsd', 'E', -1), ('distributedUsd', 'F', 1), ('netUsd', 'G', 1)]:
        actual = values['_IRR Engine'][column+row]
        same(actual*sign if isinstance(actual, (int, float)) else actual, flow[key], flow['ref']+' '+key)
same(sum(flow['calledUsd'] for flow in summary['monthlyInvestorCashFlows']), summary['headline']['capitalCalledUsd']['value'], 'Monthly capital total')
same(sum(flow['distributedUsd'] for flow in summary['monthlyInvestorCashFlows']), summary['headline']['netDistributionsUsd']['value'], 'Monthly distribution total')
same(summary['sourcesAndUses']['reserveTargetUsd']['value']-summary['sourcesAndUses']['yearOneEndingCashUsd']['value'],
     summary['sourcesAndUses']['reserveShortfallUsd']['value'], 'Reserve target shortfall')
same(summary['ownerBenefits']['annualNightsFullPortfolio']['value']*100000/summary['activeScenario']['fullCommitmentUsd']['value'],
     summary['ownerBenefits']['annualNightsPer100k']['value'], 'Pro rata owner nights')
terms = json.loads((root / 'content/investor-terms.json').read_text(encoding='utf8'))
for item in properties:
    match = re.fullmatch(r'Properties!T(\d+):Y\1', item['source']['operatingCells'])
    assert match, item['source']['operatingCells']
    row = match[1]
    same(values['Properties']['X'+row], item['stabilizedAnnualNoiUsd'], item['token']+' NOI')
    same(values['Properties']['Y'+row], item['stabilizedAnnualRentalRevenueUsd'], item['token']+' rent')
    same(values['Properties']['U'+row]-values['Properties']['T'+row], item['rentalMonths'], item['token']+' rental months')
    if terms['propertyStatus'][item['token']]['value'].startswith('Closed'):
        assert item['status'].startswith('Closed'), 'Do not regress approved closed-property status from workbook prose'

ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
with zipfile.ZipFile(args.source) as archive:
    assert not any(name.startswith('xl/externalLinks/') for name in archive.namelist()), 'External workbook links require review'
    for name in archive.namelist():
        if not re.fullmatch(r'xl/worksheets/sheet\d+\.xml', name):
            continue
        for cell in ET.fromstring(archive.read(name)).findall('.//m:sheetData/m:row/m:c', ns):
            assert cell.get('t') != 'e', f'Cached Excel error: {name}:{cell.get("r")}'
            if cell.find('m:f', ns) is not None and cell.get('t') != 'str':
                value = cell.find('m:v', ns)
                assert value is not None and value.text is not None, f'Missing numeric formula cache: {name}:{cell.get("r")}'
            formula = cell.find('m:f', ns)
            assert formula is None or '#REF!' not in (formula.text or ''), f'Broken formula reference: {name}:{cell.get("r")}'
book.close()
formulas.close()
assert hashlib.sha256(args.source.read_bytes()).hexdigest() == source_hash, 'Workbook changed during verification'
print(f'Verified {financial_count} financial source cells/formulas and {summary_count} model metrics; no Excel errors, missing numeric formula caches or external links. Source unchanged.')
